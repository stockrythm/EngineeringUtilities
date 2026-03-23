#!/usr/bin/env python3
"""
csv-import.py — Integration COE API Catalog  v2.0
===================================================
Imports the multi-sheet Excel data collection template into api-data.json.
Accepts CSVs exported from each sheet, or reads the .xlsx directly.

Usage
-----
  # From CSV exports (one per sheet):
  python3 csv-import.py apis.csv \
      --deps       dependencies.csv \
      --policies   policies.csv \
      --migration  migration.csv \
      --coe        coe.csv

  # From xlsx directly:
  python3 csv-import.py --xlsx api-catalog-template.xlsx

  # Merge new rows into existing catalog without overwriting enriched APIs:
  python3 csv-import.py apis.csv --deps dependencies.csv ... \
      --merge existing-api-data.json

  # Interactive — prompts for all paths:
  python3 csv-import.py

Sheet → catalog field mapping
------------------------------
  Sheet 1 — APIs:
    API Name              → name (also JOIN KEY for all sheets)
    Repo / JAR Name       → repoApiName
    Version               → version
    Runtime Version       → runtime  (normalised to "MuleSoft X.Y")
    Runtime Type(s)       → runtimeType[]  (split on comma)
    Java Version          → javaVersion
    Layer                 → layer
    Protocol              → protocol
    Deployment Model      → deploymentModel
    Tribe                 → tribeId  (split "tribe-id:Name" on ":")
    Squad                 → owningSquad
    Application Context   → applicationContext[]  (split on comma)
    Environments          → environments[]  (split on comma)
    Domain (L0)           → businessCapability.l0
    Capability (L1)       → businessCapability.l1
    Sub-Capability (L2)   → businessCapability.l2
    Status                → status  (normalised)
    Tags                  → tags[]  (split on comma)
    Operational Notes     → operationalNotes[]
    Confluence Link       → confluenceLinks[]
    RAML / OAS Spec URL   → ramlSpec

  Sheet 2 — Dependencies:
    API Name + Direction=CONSUMER  → consumers[]
    API Name + Direction=UPSTREAM  → upstreamSystems[]

  Sheet 3 — Policies:
    API Name + Type=GATEWAY POLICY → dependencies.gatewayPolicies[]
    API Name + Type=SHARED LIBRARY → dependencies.sharedLibraries[]

  Sheet 4 — Migration:
    All columns → migrationPlan{} + migrationStatus + migrationComplexity

  Sheet 5 — CoE:
    Columns 2-9   → security{}
    Columns 10-15 → solutionDesign{}
    Columns 16-21 → operations{}
"""

import csv
import json
import os
import re
import sys
import argparse
from datetime import datetime

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ── Helpers ───────────────────────────────────────────────────────────────────

def norm_header(h):
    return re.sub(r'\s+', ' ', str(h or '').strip().lower()).rstrip('*').strip()

def split_csv_cell(val, sep=','):
    if not val or not str(val).strip():
        return []
    return [v.strip() for v in re.split(r'[,;]+', str(val)) if v.strip()]

def norm_runtime(raw):
    if not raw or not str(raw).strip():
        return ''
    r = str(raw).strip()
    if re.match(r'^MuleSoft \d', r):
        return r
    m = re.search(r'(\d+\.\d+(?:\.\d+)?)', r)
    return f'MuleSoft {m.group(1)}' if m else r

def norm_status(raw):
    m = {'active':'Active','live':'Active','production':'Active','prod':'Active',
         'deprecated':'Deprecated','decommissioned':'Deprecated','retired':'Deprecated',
         'inactive':'Inactive','disabled':'Inactive','review':'Review','draft':'Review'}
    return m.get(str(raw or '').strip().lower(), 'Active')

def norm_envs(raw):
    parts = split_csv_cell(str(raw or ''), sep='[,;/|]')
    result = []
    for p in parts:
        p = p.upper()
        p = re.sub(r'^DEV(ELOPMENT)?$', 'DEV', p)
        p = re.sub(r'^(SYSTEM\s*INTEGRATION\s*TEST)$', 'SIT', p)
        p = re.sub(r'^(USER\s*ACCEPTANCE(\s*TEST)?)$', 'UAT', p)
        p = re.sub(r'^PROD(UCTION)?$', 'PROD', p)
        if p:
            result.append(p)
    return sorted(set(result))

def infer_java(runtime):
    m = re.search(r'(\d+)\.(\d+)', str(runtime or ''))
    if m:
        major, minor = int(m.group(1)), int(m.group(2))
        if major == 4 and minor <= 6:
            return 'Java 8'
        if major == 4 and minor >= 7:
            return 'Java 17'
        if major == 3:
            return 'Java 8'
    return ''

def make_id():
    now = datetime.now()
    return (f"api-{now.year}{str(now.month).zfill(2)}{str(now.day).zfill(2)}"
            f"{str(now.hour).zfill(2)}{str(now.minute).zfill(2)}{str(now.second).zfill(2)}"
            f"{str(now.microsecond)[:3]}")

def find_col(headers, *variants):
    for h in headers:
        for v in variants:
            if h == v or h.startswith(v):
                return headers.index(h)
    return -1

def get(row, idx, default=''):
    if idx < 0 or idx >= len(row):
        return default
    return str(row[idx] or '').strip()

# ── Skeleton ──────────────────────────────────────────────────────────────────

def skeleton():
    return {
        "id": "", "name": "", "repoApiName": "", "version": "",
        "layer": "", "protocol": "REST",
        "runtime": "", "runtimeType": [], "javaVersion": "", "deploymentModel": "",
        "environments": [], "tribeId": "", "owningSquad": "",
        "applicationContext": [],
        "businessCapability": {"l0": "", "l1": "", "l2": ""},
        "consumers": [], "upstreamSystems": [],
        "status": "Active", "migrationStatus": "Not Started",
        "migrationComplexity": "Medium",
        "ramlSpec": "", "confluenceLinks": [], "designReferences": [],
        "operationalNotes": [], "tags": [],
        "lastReviewedBy": "", "lastReviewedDate": "",
        "security": {
            "authScheme": "", "oauthScopes": [], "tlsVersion": "",
            "gatewayPolicies": [], "dataClassification": "",
            "complianceFlags": [], "secretsVaultPath": "",
            "certExpiry": "", "knownVulnerabilities": []
        },
        "solutionDesign": {
            "architecturalPattern": "", "versioningStrategy": "",
            "errorHandlingStrategy": "", "slaTarget": "", "asyncApiSpec": "",
            "adrLinks": [], "technicalDebt": [], "domainEvents": []
        },
        "operations": {
            "oncallContact": "", "monitoringLinks": [], "alertingThreshold": "",
            "maintenanceWindow": "", "runbookLinks": [], "incidentRefs": [],
            "knownFragilities": [], "pipelineLink": "", "uptimeSla": ""
        },
        "dependencies": {"gatewayPolicies": [], "sharedLibraries": []},
        "migrationPlan": {
            "wave": "unassigned", "targetStartDate": "", "targetDate": "",
            "completionDate": "", "blockerReason": "", "blockerOwner": ""
        }
    }

# ── Sheet readers ─────────────────────────────────────────────────────────────

def read_csv_file(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.reader(f))
    if not rows:
        return [], []
    headers = [norm_header(h) for h in rows[0]]
    return headers, rows[1:]

def read_xlsx_sheet(wb, sheet_name):
    if sheet_name not in wb.sheetnames:
        return [], []
    ws = wb[sheet_name]
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    if not rows:
        return [], []
    # Find header row — look for row where a cell is exactly "api name" or "api name *"
    hdr_row = 0
    for i, row in enumerate(rows):
        norm = [norm_header(str(c or '')) for c in row]
        if any(h in ('api name', 'api name *', 'api name*') for h in norm):
            hdr_row = i
            break
    headers = [norm_header(str(c or '')) for c in rows[hdr_row]]
    data = rows[hdr_row + 1:]
    return headers, [[str(c or '').strip() if c is not None else '' for c in row] for row in data]

# ── Sheet 1 — APIs ────────────────────────────────────────────────────────────

def parse_apis(headers, rows, existing_ids, warnings):
    apis = {}
    ci = {
        'name':     find_col(headers, 'api name'),
        'repo':     find_col(headers, 'repo / jar name', 'repo/jar name', 'repo name', 'jar name', 'jar version'),
        'version':  find_col(headers, 'version', 'api version'),
        'runtime':  find_col(headers, 'runtime version', 'runtime'),
        'rttype':   find_col(headers, 'runtime type', 'runtime types'),
        'java':     find_col(headers, 'java version', 'java'),
        'layer':    find_col(headers, 'layer', 'type', 'api type'),
        'protocol': find_col(headers, 'protocol'),
        'deploy':   find_col(headers, 'deployment model', 'deployment'),
        'tribe':    find_col(headers, 'tribe'),
        'squad':    find_col(headers, 'squad'),
        'appctx':   find_col(headers, 'application context'),
        'envs':     find_col(headers, 'environments', 'environment'),
        'l0':       find_col(headers, 'domain (l0)', 'domain', 'l0'),
        'l1':       find_col(headers, 'capability (l1)', 'capability', 'l1'),
        'l2':       find_col(headers, 'sub-capability (l2)', 'sub-capability', 'l2'),
        'status':   find_col(headers, 'status', 'api status'),
        'tags':     find_col(headers, 'tags'),
        'notes':    find_col(headers, 'operational notes'),
        'conf':     find_col(headers, 'confluence link', 'confluence'),
        'raml':     find_col(headers, 'raml / oas spec url', 'raml', 'spec url', 'oas'),
        'hotspot':  find_col(headers, 'hotspot api', 'hotspot'),
        'inscope':  find_col(headers, 'in scope of targeted projects', 'in scope'),
        'project':  find_col(headers, 'project name', 'project'),
        'eymaster': find_col(headers, 'included in ey master', 'ey master'),
    }

    for r_num, row in enumerate(rows, 2):
        if not any(str(c or '').strip() for c in row):
            continue
        name = get(row, ci['name'])
        if not name:
            continue

        api = skeleton()
        api['id'] = make_id()
        while api['id'] in existing_ids:
            import time; time.sleep(0.001); api['id'] = make_id()
        existing_ids.add(api['id'])

        api['name']          = name
        api['repoApiName']   = get(row, ci['repo']) or name
        api['version']       = get(row, ci['version'])
        api['runtime']       = norm_runtime(get(row, ci['runtime']))
        api['runtimeType']   = split_csv_cell(get(row, ci['rttype']))
        api['javaVersion']   = get(row, ci['java']) or infer_java(api['runtime'])
        api['layer']         = get(row, ci['layer']).upper() if get(row, ci['layer']) else ''
        api['protocol']      = get(row, ci['protocol']) or 'REST'
        api['deploymentModel'] = get(row, ci['deploy'])
        api['owningSquad']   = get(row, ci['squad'])
        api['applicationContext'] = split_csv_cell(get(row, ci['appctx']))
        api['environments']  = norm_envs(get(row, ci['envs']))
        api['status']        = norm_status(get(row, ci['status']))
        api['tags']          = split_csv_cell(get(row, ci['tags']))
        api['businessCapability'] = {
            'l0': get(row, ci['l0']),
            'l1': get(row, ci['l1']),
            'l2': get(row, ci['l2']),
        }

        # Tribe: split "tribe-id:Tribe Name" on ":"
        tribe_raw = get(row, ci['tribe'])
        if ':' in tribe_raw:
            api['tribeId'] = tribe_raw.split(':')[0].strip()
        elif tribe_raw:
            api['tribeId'] = tribe_raw

        # Notes / docs
        notes = get(row, ci['notes'])
        if notes:
            api['operationalNotes'].append(notes)
        project = get(row, ci['project'])
        if project:
            api['operationalNotes'].append(f'Project: {project}')
        conf = get(row, ci['conf'])
        if conf:
            api['confluenceLinks'].append({'label': 'Confluence', 'url': conf})
        api['ramlSpec'] = get(row, ci['raml'])

        # Legacy CSV columns (hotspot → complexity, in scope → migStatus)
        hotspot = get(row, ci['hotspot']).lower()
        if hotspot in ('yes','y','true','1','x'):
            api['migrationComplexity'] = 'High'
        inscope = get(row, ci['inscope']).lower()
        if inscope in ('yes','y','true','1','x'):
            api['migrationStatus'] = 'In Analysis'
        eymaster = get(row, ci['eymaster']).lower()
        if eymaster in ('yes','y','true','1','x'):
            if 'EY-Master' not in api['tags']:
                api['tags'].append('EY-Master')

        apis[name.lower()] = api

    return apis

# ── Sheet 2 — Dependencies ────────────────────────────────────────────────────

def parse_deps(headers, rows, apis_by_name, warnings):
    ci = {
        'name':     find_col(headers, 'api name'),
        'direction':find_col(headers, 'direction'),
        'contract': find_col(headers, 'contract / protocol', 'contract/protocol', 'protocol', 'contract type'),
        'system':   find_col(headers, 'connected system', 'system'),
        'envs':     find_col(headers, 'environments', 'environment'),
        'notes':    find_col(headers, 'notes'),
    }
    for r_num, row in enumerate(rows, 2):
        if not any(str(c or '').strip() for c in row):
            continue
        name = get(row, ci['name'])
        direction = get(row, ci['direction']).upper()
        system = get(row, ci['system'])
        if not name or not system:
            continue
        key = name.lower()
        if key not in apis_by_name:
            warnings.append(f'  Dependencies row {r_num}: API "{name}" not found in APIs sheet — skipped')
            continue
        api = apis_by_name[key]
        envs = norm_envs(get(row, ci['envs']))
        notes = get(row, ci['notes'])
        contract = get(row, ci['contract'])

        if direction == 'CONSUMER':
            api['consumers'].append({
                'system': system, 'environments': envs,
                'contractType': contract, 'notes': notes
            })
        elif direction == 'UPSTREAM':
            api['upstreamSystems'].append({
                'system': system, 'environments': envs,
                'protocol': contract, 'notes': notes
            })
        else:
            warnings.append(f'  Dependencies row {r_num}: Unknown direction "{direction}" — expected CONSUMER or UPSTREAM')

# ── Sheet 3 — Policies ────────────────────────────────────────────────────────

def parse_policies(headers, rows, apis_by_name, warnings):
    ci = {
        'name':    find_col(headers, 'api name'),
        'type':    find_col(headers, 'type'),
        'polname': find_col(headers, 'name'),
        'version': find_col(headers, 'version'),
        'scope':   find_col(headers, 'scope / group id', 'scope/group id', 'scope', 'group id'),
        'exurl':   find_col(headers, 'exchange url', 'exchange'),
        'notes':   find_col(headers, 'notes'),
    }
    for r_num, row in enumerate(rows, 2):
        if not any(str(c or '').strip() for c in row):
            continue
        name    = get(row, ci['name'])
        poltype = get(row, ci['type']).upper()
        polname = get(row, ci['polname'])
        if not name or not polname:
            continue
        key = name.lower()
        if key not in apis_by_name:
            warnings.append(f'  Policies row {r_num}: API "{name}" not found — skipped')
            continue
        api   = apis_by_name[key]
        ver   = get(row, ci['version'])
        scope = get(row, ci['scope'])
        exurl = get(row, ci['exurl'])
        notes = get(row, ci['notes'])

        if poltype == 'GATEWAY POLICY':
            api['dependencies']['gatewayPolicies'].append({
                'name': polname, 'version': ver, 'scope': scope or 'API Manager', 'notes': notes
            })
        elif poltype == 'SHARED LIBRARY':
            api['dependencies']['sharedLibraries'].append({
                'name': polname, 'groupId': scope, 'version': ver,
                'scope': 'POM dependency', 'exchangeUrl': exurl
            })
        else:
            warnings.append(f'  Policies row {r_num}: Unknown type "{poltype}" — expected GATEWAY POLICY or SHARED LIBRARY')

# ── Sheet 4 — Migration ───────────────────────────────────────────────────────

def parse_migration(headers, rows, apis_by_name, warnings):
    ci = {
        'name':       find_col(headers, 'api name'),
        'wave':       find_col(headers, 'wave'),
        'migstatus':  find_col(headers, 'migration status'),
        'complexity': find_col(headers, 'complexity'),
        'startdate':  find_col(headers, 'target start date', 'start date'),
        'enddate':    find_col(headers, 'target end date', 'end date', 'target date'),
        'actual':     find_col(headers, 'actual completion', 'completion date'),
        'blocker':    find_col(headers, 'blocker reason'),
        'owner':      find_col(headers, 'blocker owner'),
    }
    wave_map = {
        'wave 1': 'wave-1', 'wave1': 'wave-1',
        'wave 2': 'wave-2', 'wave2': 'wave-2',
        'wave 3': 'wave-3', 'wave3': 'wave-3',
        'wave 4': 'wave-4', 'wave4': 'wave-4',
        'wave 5': 'wave-5', 'wave5': 'wave-5',
        'unassigned': 'unassigned', '': 'unassigned',
    }
    for r_num, row in enumerate(rows, 2):
        if not any(str(c or '').strip() for c in row):
            continue
        name = get(row, ci['name'])
        if not name:
            continue
        key = name.lower()
        if key not in apis_by_name:
            warnings.append(f'  Migration row {r_num}: API "{name}" not found — skipped')
            continue
        api = apis_by_name[key]
        wave_raw = get(row, ci['wave']).lower()
        api['migrationStatus']   = get(row, ci['migstatus']) or 'Not Started'
        api['migrationComplexity'] = get(row, ci['complexity']) or 'Medium'
        api['migrationPlan'] = {
            'wave':            wave_map.get(wave_raw, 'unassigned'),
            'targetStartDate': get(row, ci['startdate']),
            'targetDate':      get(row, ci['enddate']),
            'completionDate':  get(row, ci['actual']),
            'blockerReason':   get(row, ci['blocker']),
            'blockerOwner':    get(row, ci['owner']),
        }

# ── Sheet 5 — CoE ─────────────────────────────────────────────────────────────

def parse_coe(headers, rows, apis_by_name, warnings):
    ci = {
        'name':      find_col(headers, 'api name'),
        'auth':      find_col(headers, 'auth scheme'),
        'scopes':    find_col(headers, 'oauth scopes'),
        'tls':       find_col(headers, 'tls version', 'tls'),
        'datacls':   find_col(headers, 'data classification'),
        'compliance':find_col(headers, 'compliance flags'),
        'certexp':   find_col(headers, 'cert expiry date', 'cert expiry'),
        'vault':     find_col(headers, 'secrets vault path', 'vault path'),
        'vulns':     find_col(headers, 'known vulnerabilities'),
        'archpat':   find_col(headers, 'architectural pattern'),
        'versioning':find_col(headers, 'versioning strategy'),
        'errhandling':find_col(headers,'error handling strategy', 'error handling'),
        'sla':       find_col(headers, 'sla target', 'sla'),
        'adrs':      find_col(headers, 'adr links', 'adr'),
        'debt':      find_col(headers, 'technical debt'),
        'oncall':    find_col(headers, 'on-call contact', 'oncall'),
        'uptime':    find_col(headers, 'uptime sla', 'uptime'),
        'alerting':  find_col(headers, 'alerting threshold'),
        'maint':     find_col(headers, 'maintenance window'),
        'monitoring':find_col(headers, 'monitoring link'),
        'fragilities':find_col(headers,'known fragilities'),
    }
    for r_num, row in enumerate(rows, 2):
        if not any(str(c or '').strip() for c in row):
            continue
        name = get(row, ci['name'])
        if not name:
            continue
        key = name.lower()
        if key not in apis_by_name:
            warnings.append(f'  CoE row {r_num}: API "{name}" not found — skipped')
            continue
        api = apis_by_name[key]
        api['security'].update({
            'authScheme':         get(row, ci['auth']),
            'oauthScopes':        split_csv_cell(get(row, ci['scopes'])),
            'tlsVersion':         get(row, ci['tls']),
            'dataClassification': get(row, ci['datacls']),
            'complianceFlags':    split_csv_cell(get(row, ci['compliance'])),
            'certExpiry':         get(row, ci['certexp']),
            'secretsVaultPath':   get(row, ci['vault']),
            'knownVulnerabilities': [get(row, ci['vulns'])] if get(row, ci['vulns']) else [],
        })
        api['solutionDesign'].update({
            'architecturalPattern':  get(row, ci['archpat']),
            'versioningStrategy':    get(row, ci['versioning']),
            'errorHandlingStrategy': get(row, ci['errhandling']),
            'slaTarget':             get(row, ci['sla']),
            'adrLinks':              split_csv_cell(get(row, ci['adrs'])),
            'technicalDebt':         [get(row, ci['debt'])] if get(row, ci['debt']) else [],
        })
        api['operations'].update({
            'oncallContact':    get(row, ci['oncall']),
            'uptimeSla':        get(row, ci['uptime']),
            'alertingThreshold':get(row, ci['alerting']),
            'maintenanceWindow':get(row, ci['maint']),
            'monitoringLinks':  [{'label':'Dashboard','url':get(row,ci['monitoring'])}]
                                if get(row, ci['monitoring']) else [],
            'knownFragilities': [get(row, ci['fragilities'])] if get(row, ci['fragilities']) else [],
        })

# ── Base catalog ──────────────────────────────────────────────────────────────

def base_catalog(meta=None):
    return {
        "meta": {
            "version": "1.6.0",
            "lastUpdated": datetime.now().strftime("%Y-%m-%d"),
            "organization": meta.get("organization","") if meta else "",
            "catalogOwner": meta.get("catalogOwner","") if meta else "",
        },
        "layerSuggestions": ["EAPI","PAPI","SAPI","BATCH","EVENT","UTILITY","FACADE"],
        "contractTypes": ["REST","SOAP","Kafka","AsyncAPI","GraphQL","SFTP","JMS","Other"],
        "runtimeVersions": ["MuleSoft 4.1","MuleSoft 4.2","MuleSoft 4.3","MuleSoft 4.4",
                            "MuleSoft 4.5","MuleSoft 4.6","MuleSoft 4.7","MuleSoft 4.8",
                            "MuleSoft 4.9","MuleSoft 3.9 (Legacy)"],
        "runtimeTypes": ["AWS Egress Runtime","AWS Ingress Runtime","IGZ Egress",
                         "IGZ Ingress","Qsystem","Process","Experience","System"],
        "deploymentModels": ["OnPrem","AWS","Azure","Hybrid","CloudHub"],
        "javaVersions": ["Java 8","Java 11","Java 17","Java 21"],
        "complianceFlagOptions": ["PCI-DSS","GDPR","SOX","HIPAA","ISO27001"],
        "dataClassifications": ["Public","Internal","Confidential","Restricted"],
        "tribes": meta.get("tribes",[]) if meta else [],
        "businessCapabilities": meta.get("businessCapabilities",[]) if meta else [],
        "migrationWaves": meta.get("migrationWaves",[
            {"id":"wave-1","name":"Wave 1","description":"Foundation & Platform APIs",
             "startDate":"","endDate":"","color":"#6366f1"},
            {"id":"wave-2","name":"Wave 2","description":"Core Business APIs",
             "startDate":"","endDate":"","color":"#3b82f6"},
            {"id":"wave-3","name":"Wave 3","description":"Experience & Channel APIs",
             "startDate":"","endDate":"","color":"#22d3ee"},
            {"id":"unassigned","name":"Unassigned","description":"",
             "startDate":"","endDate":"","color":"#64748b"},
        ]) if meta else [],
        "apis": []
    }

# ── Main ──────────────────────────────────────────────────────────────────────

def run(apis_path, deps_path=None, policies_path=None, migration_path=None,
        coe_path=None, xlsx_path=None, output_path='api-data.json', merge_path=None):

    warnings = []
    skipped  = []
    dupes    = []

    existing_catalog = None
    existing_names   = set()
    existing_ids     = set()
    if merge_path:
        with open(merge_path) as f:
            existing_catalog = json.load(f)
        existing_names = {a['name'].lower() for a in existing_catalog.get('apis',[])}
        existing_ids   = {a['id'] for a in existing_catalog.get('apis',[])}
        print(f'Merge mode: {len(existing_names)} existing APIs loaded from {merge_path}')

    # ── Read sheets ───────────────────────────────────────────
    if xlsx_path:
        if not HAS_OPENPYXL:
            print('ERROR: openpyxl required for xlsx mode. Run: pip install openpyxl')
            sys.exit(1)
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        h1,r1 = read_xlsx_sheet(wb, 'APIs')
        h2,r2 = read_xlsx_sheet(wb, 'Dependencies')
        h3,r3 = read_xlsx_sheet(wb, 'Policies')
        h4,r4 = read_xlsx_sheet(wb, 'Migration')
        h5,r5 = read_xlsx_sheet(wb, 'CoE')
    else:
        h1,r1 = read_csv_file(apis_path)
        h2,r2 = read_csv_file(deps_path)       if deps_path       else ([],[])
        h3,r3 = read_csv_file(policies_path)   if policies_path   else ([],[])
        h4,r4 = read_csv_file(migration_path)  if migration_path  else ([],[])
        h5,r5 = read_csv_file(coe_path)        if coe_path        else ([],[])

    # ── Parse APIs ────────────────────────────────────────────
    apis_by_name = parse_apis(h1, r1, existing_ids, warnings)

    # ── Filter duplicates for merge ───────────────────────────
    if merge_path:
        for name_key in list(apis_by_name.keys()):
            if name_key in existing_names:
                dupes.append(apis_by_name[name_key]['name'])
                del apis_by_name[name_key]

    # ── Enrich from other sheets ──────────────────────────────
    if r2: parse_deps(h2, r2, apis_by_name, warnings)
    if r3: parse_policies(h3, r3, apis_by_name, warnings)
    if r4: parse_migration(h4, r4, apis_by_name, warnings)
    if r5: parse_coe(h5, r5, apis_by_name, warnings)

    imported = list(apis_by_name.values())

    # ── Build catalog ─────────────────────────────────────────
    meta_src = existing_catalog if merge_path and existing_catalog else None
    if merge_path and existing_catalog:
        catalog = existing_catalog
        catalog['apis'].extend(imported)
        catalog['meta']['lastUpdated'] = datetime.now().strftime("%Y-%m-%d")
    else:
        catalog = base_catalog(meta_src)
        catalog['apis'] = imported

    with open(output_path,'w',encoding='utf-8') as f:
        json.dump(catalog, f, indent=2, ensure_ascii=False)

    # ── Report ────────────────────────────────────────────────
    report_path = output_path.replace('.json','-import-report.txt')
    lines = [
        '═'*62, '  Integration COE — Multi-Sheet Import Report',
        f'  {datetime.now().strftime("%Y-%m-%d %H:%M")}', '═'*62, '',
        f'  Mode       : {"Merge into "+merge_path if merge_path else "Fresh import"}',
        f'  Output     : {output_path}','',
        '─'*62,'  SUMMARY','─'*62,
        f'  ✅ APIs imported      : {len(imported)}',
        f'  🔁 Duplicates skipped : {len(dupes)}',
        f'  ⚠  Warnings           : {len(warnings)}',
        '',
        f'  Consumers populated  : {sum(len(a["consumers"]) for a in imported)}',
        f'  Upstream populated   : {sum(len(a["upstreamSystems"]) for a in imported)}',
        f'  Policies populated   : {sum(len(a["dependencies"]["gatewayPolicies"]) for a in imported)}',
        f'  Libraries populated  : {sum(len(a["dependencies"]["sharedLibraries"]) for a in imported)}',
        f'  With migration plan  : {sum(1 for a in imported if a["migrationPlan"]["wave"]!="unassigned")}',
        f'  With CoE security    : {sum(1 for a in imported if a["security"]["authScheme"])}',
        '',
    ]
    if dupes:
        lines += ['─'*62,'  DUPLICATES (not imported)','─'*62]
        for n in dupes: lines.append(f'  • {n}')
        lines.append('')
    if warnings:
        lines += ['─'*62,'  WARNINGS','─'*62]
        lines.extend(warnings); lines.append('')
    lines += [
        '─'*62,'  ENRICHMENT CHECKLIST (post-import)','─'*62,
        '  1. tribeId          → assign via editor.html (required for maps)',
        '  2. consumers[]      → verify system names resolve in catalog',
        '  3. migrationPlan    → assign waves in migration.html planner',
        '  4. CoE blocks       → run workshops per tribe',
        '  5. applicationContext[] → confirm multi-app APIs',
        '═'*62,
    ]
    with open(report_path,'w',encoding='utf-8') as f:
        f.write('\n'.join(lines))

    print()
    print('═'*62)
    print(f'  ✅  Imported   : {len(imported)} APIs')
    if dupes:    print(f'  🔁  Duplicates : {len(dupes)} skipped')
    if warnings: print(f'  ⚠   Warnings   : {len(warnings)} — see report')
    print(f'\n  Output JSON  → {output_path}')
    print(f'  Report       → {report_path}')
    print('═'*62)
    print()

# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Import Excel/CSV catalog data to api-data.json')
    parser.add_argument('apis_csv',      nargs='?',  help='Sheet 1 CSV (APIs)')
    parser.add_argument('--deps',        metavar='CSV', help='Sheet 2 CSV (Dependencies)')
    parser.add_argument('--policies',    metavar='CSV', help='Sheet 3 CSV (Policies)')
    parser.add_argument('--migration',   metavar='CSV', help='Sheet 4 CSV (Migration)')
    parser.add_argument('--coe',         metavar='CSV', help='Sheet 5 CSV (CoE Metadata)')
    parser.add_argument('--xlsx',        metavar='XLSX',help='Read all sheets from .xlsx directly')
    parser.add_argument('--output', '-o',metavar='JSON',default='api-data.json')
    parser.add_argument('--merge',       metavar='JSON',help='Merge into existing catalog')
    args = parser.parse_args()

    if not args.apis_csv and not args.xlsx:
        print('\nIntegration COE — Multi-Sheet Import Tool')
        print('─'*44)
        mode = input('  Mode? [1] CSV files  [2] XLSX direct  : ').strip()
        if mode == '2':
            xlsx = input('  Path to .xlsx file: ').strip().strip('"')
            out  = input('  Output [api-data.json]: ').strip() or 'api-data.json'
            mg   = input('  Merge into existing JSON? (path or blank): ').strip().strip('"') or None
            run(None, xlsx_path=xlsx, output_path=out, merge_path=mg or None)
        else:
            apis = input('  Sheet 1 — APIs CSV: ').strip().strip('"')
            deps = input('  Sheet 2 — Dependencies CSV (blank to skip): ').strip().strip('"') or None
            pols = input('  Sheet 3 — Policies CSV (blank to skip): ').strip().strip('"') or None
            mig  = input('  Sheet 4 — Migration CSV (blank to skip): ').strip().strip('"') or None
            coe  = input('  Sheet 5 — CoE CSV (blank to skip): ').strip().strip('"') or None
            out  = input('  Output [api-data.json]: ').strip() or 'api-data.json'
            mg   = input('  Merge into existing JSON? (path or blank): ').strip().strip('"') or None
            run(apis, deps, pols, mig, coe, output_path=out, merge_path=mg or None)
    else:
        if args.xlsx:
            run(None, xlsx_path=args.xlsx, output_path=args.output, merge_path=args.merge)
        else:
            run(args.apis_csv, args.deps, args.policies, args.migration, args.coe,
                output_path=args.output, merge_path=args.merge)
